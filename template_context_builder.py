"""
template_context_builder.py

Shared helpers for deriving presentation context from a source PowerPoint
template.
"""

import os
from collections import Counter

from pptx import Presentation
from pptx.util import Emu

DEFAULT_CONTEXT = {
    "source_template": None,
    "slide_width_in": 13.333,
    "slide_height_in": 7.5,
    "fonts": {"title": "Cambria", "body": "Calibri"},
    "colors": {
        "primary": "1E2761",
        "secondary": "CADCFC",
        "accent": "FFFFFF",
        "background": "FFFFFF",
        "text": "212121",
    },
    "images": [],
    "summary": "No template supplied — using a default 16:9 'Midnight Executive' style context.",
}


def _emu_to_in(value):
    return round(Emu(value).inches, 3) if value is not None else None


def _collect_fonts(prs):
    counts = Counter()
    title_counts = Counter()

    def walk(shapes, is_title_context=False):
        for shape in shapes:
            if not shape.has_text_frame:
                continue
            is_title = is_title_context or (
                shape.is_placeholder and shape.placeholder_format.idx == 0
            )
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        if is_title:
                            title_counts[run.font.name] += 1
                        else:
                            counts[run.font.name] += 1

    for master in prs.slide_masters:
        walk(master.shapes)
        for layout in master.slide_layouts:
            walk(layout.shapes)
    for slide in prs.slides:
        walk(slide.shapes)

    title_font = title_counts.most_common(1)[0][0] if title_counts else None
    body_font = counts.most_common(1)[0][0] if counts else None
    return title_font, body_font


def _collect_theme_colors(prs):
    colors = {}
    try:
        master = prs.slide_masters[0]
        theme_part = master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
        )
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = theme_part._element.find(".//a:clrScheme", ns)
        if scheme is not None:
            for child in scheme:
                tag = child.tag.split("}")[-1]
                srgb = child.find("a:srgbClr", ns)
                sys_clr = child.find("a:sysClr", ns)
                if srgb is not None:
                    colors[tag] = srgb.get("val")
                elif sys_clr is not None:
                    colors[tag] = sys_clr.get("lastClr")
    except Exception:
        pass
    return colors


def _extract_images(prs, output_dir):
    assets_dir = os.path.join(output_dir, "assets")
    os.makedirs(assets_dir, exist_ok=True)
    saved = []
    seen_hashes = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13 or getattr(shape, "image", None):
                try:
                    image = shape.image
                except Exception:
                    continue
                if image.sha1 in seen_hashes:
                    continue
                seen_hashes.add(image.sha1)
                fname = f"asset_{len(saved) + 1}.{image.ext}"
                fpath = os.path.join(assets_dir, fname)
                with open(fpath, "wb") as f:
                    f.write(image.blob)
                saved.append(os.path.relpath(fpath, output_dir))
    return saved


def build_context(template_path, output_dir):
    if not template_path:
        return dict(DEFAULT_CONTEXT)

    prs = Presentation(template_path)
    title_font, body_font = _collect_fonts(prs)
    theme_colors = _collect_theme_colors(prs)
    images = _extract_images(prs, output_dir)

    return {
        "source_template": os.path.basename(template_path),
        "slide_width_in": _emu_to_in(prs.slide_width),
        "slide_height_in": _emu_to_in(prs.slide_height),
        "fonts": {
            "title": title_font or DEFAULT_CONTEXT["fonts"]["title"],
            "body": body_font or DEFAULT_CONTEXT["fonts"]["body"],
        },
        "colors": {
            "primary": theme_colors.get("dk2")
            or theme_colors.get("accent1")
            or DEFAULT_CONTEXT["colors"]["primary"],
            "secondary": theme_colors.get("accent2")
            or DEFAULT_CONTEXT["colors"]["secondary"],
            "accent": theme_colors.get("accent3")
            or DEFAULT_CONTEXT["colors"]["accent"],
            "background": theme_colors.get("lt1")
            or DEFAULT_CONTEXT["colors"]["background"],
            "text": theme_colors.get("dk1") or DEFAULT_CONTEXT["colors"]["text"],
        },
        "images": images,
        "summary": (
            f"Template '{os.path.basename(template_path)}': "
            f"{_emu_to_in(prs.slide_width)}x{_emu_to_in(prs.slide_height)} in, "
            f"title font '{title_font}', body font '{body_font}', "
            f"{len(images)} embedded image asset(s) extracted."
        ),
    }
