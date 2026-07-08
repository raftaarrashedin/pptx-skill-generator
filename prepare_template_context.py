"""
prepare_template_context.py

Reads a source .pptx template and extracts the minimum design context needed
to generate slides that look like they belong to it: slide size, fonts,
brand-ish colors, and any embedded image/logo assets.

If no template is given, writes a sensible default context (16:9, safe-font
pairing, neutral palette) so the rest of the pipeline still has something to
work with.

Output: output/template_context.json

Usage:
    python prepare_template_context.py --template path/to/template.pptx
    python prepare_template_context.py                 # no template -> defaults
"""

import argparse
import json
import os
from collections import Counter
from datetime import datetime, timezone

from pptx import Presentation
from pptx.util import Emu

from extract_presentation_template import extract_template_package
from template_actor import (
    package_dir_for_template,
    package_exists,
    package_has_svgs,
    stage_package_into_output,
)
from template_catalog import list_templates

EMU_PER_INCH = 914400

DEFAULT_CONTEXT = {
    "source_template": None,
    "slide_width_in": 13.333,
    "slide_height_in": 7.5,
    "fonts": {"title": "Cambria", "body": "Calibri"},
    "colors": {
        "primary": "1E2761",
        "secondary": "CADCFC",
        "accent": "FFFFFF",
        "background": "FFFFFF",
        "text": "212121",
    },
    "images": [],
    "summary": "No template supplied — using a default 16:9 'Midnight Executive' style context.",
}


def log(output_dir, message):
    path = os.path.join(output_dir, "run.log")
    with open(path, "a", encoding="utf-8") as f:
        f.write(
            f"{datetime.now(timezone.utc).isoformat()} | prepare_template_context | {message}\n"
        )


def _emu_to_in(value):
    return round(Emu(value).inches, 3) if value is not None else None


def _collect_fonts(prs):
    """Walk slide masters + layouts + slides, tally font names used in runs."""
    counts = Counter()
    title_counts = Counter()

    def walk(shapes, is_title_context=False):
        for shape in shapes:
            if not shape.has_text_frame:
                continue
            is_title = is_title_context or (
                shape.is_placeholder and shape.placeholder_format.idx == 0
            )
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        if is_title:
                            title_counts[run.font.name] += 1
                        else:
                            counts[run.font.name] += 1

    for master in prs.slide_masters:
        walk(master.shapes)
        for layout in master.slide_layouts:
            walk(layout.shapes)
    for slide in prs.slides:
        walk(slide.shapes)

    title_font = title_counts.most_common(1)[0][0] if title_counts else None
    body_font = counts.most_common(1)[0][0] if counts else None
    return title_font, body_font


def _collect_theme_colors(prs):
    """Pull the theme color scheme (accent1/2, dk1, lt1) straight from master XML."""
    colors = {}
    try:
        master = prs.slide_masters[0]
        theme = master.element.getroottree().getroot()
        # python-pptx doesn't expose theme directly off the slide master object
        # in a friendly way across versions, so we reach into the part relationship.
        theme_part = master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
        )
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = theme_part._element.find(".//a:clrScheme", ns)
        if scheme is not None:
            for child in scheme:
                tag = child.tag.split("}")[-1]
                srgb = child.find("a:srgbClr", ns)
                sys_clr = child.find("a:sysClr", ns)
                if srgb is not None:
                    colors[tag] = srgb.get("val")
                elif sys_clr is not None:
                    colors[tag] = sys_clr.get("lastClr")
    except Exception:
        pass
    return colors


def _extract_images(prs, output_dir):
    """Save embedded picture assets so later steps can reference real logos/images."""
    assets_dir = os.path.join(output_dir, "assets")
    os.makedirs(assets_dir, exist_ok=True)
    saved = []
    seen_hashes = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13 or getattr(shape, "image", None):  # PICTURE
                try:
                    image = shape.image
                except Exception:
                    continue
                if image.sha1 in seen_hashes:
                    continue
                seen_hashes.add(image.sha1)
                fname = f"asset_{len(saved) + 1}.{image.ext}"
                fpath = os.path.join(assets_dir, fname)
                with open(fpath, "wb") as f:
                    f.write(image.blob)
                saved.append(os.path.relpath(fpath, output_dir))
    return saved


def build_context(template_path, output_dir):
    if not template_path:
        ctx = dict(DEFAULT_CONTEXT)
        return ctx

    prs = Presentation(template_path)
    title_font, body_font = _collect_fonts(prs)
    theme_colors = _collect_theme_colors(prs)
    images = _extract_images(prs, output_dir)

    ctx = {
        "source_template": os.path.basename(template_path),
        "slide_width_in": _emu_to_in(prs.slide_width),
        "slide_height_in": _emu_to_in(prs.slide_height),
        "fonts": {
            "title": title_font or DEFAULT_CONTEXT["fonts"]["title"],
            "body": body_font or DEFAULT_CONTEXT["fonts"]["body"],
        },
        "colors": {
            "primary": theme_colors.get("dk2")
            or theme_colors.get("accent1")
            or DEFAULT_CONTEXT["colors"]["primary"],
            "secondary": theme_colors.get("accent2")
            or DEFAULT_CONTEXT["colors"]["secondary"],
            "accent": theme_colors.get("accent3")
            or DEFAULT_CONTEXT["colors"]["accent"],
            "background": theme_colors.get("lt1")
            or DEFAULT_CONTEXT["colors"]["background"],
            "text": theme_colors.get("dk1") or DEFAULT_CONTEXT["colors"]["text"],
        },
        "images": images,
        "summary": (
            f"Template '{os.path.basename(template_path)}': "
            f"{_emu_to_in(prs.slide_width)}x{_emu_to_in(prs.slide_height)} in, "
            f"title font '{title_font}', body font '{body_font}', "
            f"{len(images)} embedded image asset(s) extracted."
        ),
    }
    return ctx


def main():
    parser = argparse.ArgumentParser(
        description="Extract design context from a pptx template."
    )
    parser.add_argument(
        "--template",
        default=None,
        help="Path to source .pptx template or a saved template name such as 'Nature Journal'",
    )
    parser.add_argument(
        "--list-templates",
        action="store_true",
        help="Print available saved template names and exit",
    )
    parser.add_argument(
        "--output-dir", default="output", help="Output directory (default: output)"
    )
    args = parser.parse_args()

    if args.list_templates:
        templates = list_templates()
        if not templates:
            print(
                "No saved templates found. Add .pptx files to the project root or templates/."
            )
            return
        print("Available templates:")
        for template in templates:
            package_dir = package_dir_for_template(template.path)
            status = (
                "cached"
                if package_exists(package_dir) and package_has_svgs(package_dir)
                else "not extracted"
            )
            print(f"- {template.name} [{status}] -> {template.path}")
        return

    os.makedirs(args.output_dir, exist_ok=True)
    log(args.output_dir, "started")
    if args.template:
        try:
            package_dir, _, reused = extract_template_package(args.template)
            ctx = stage_package_into_output(package_dir, args.output_dir)
            ctx["summary"] = (
                f"{ctx.get('summary', '')} "
                f"[template package: {'reused' if reused else 'newly extracted'} from {ctx.get('source_template_name') or ctx.get('source_template')}]"
            ).strip()
            log(args.output_dir, f"template package used: {package_dir}")
        except Exception as e:
            log(args.output_dir, f"ERROR: {e}")
            raise
    else:
        log(args.output_dir, "no template provided, using defaults")
        try:
            ctx = build_context(None, args.output_dir)
        except Exception as e:
            log(args.output_dir, f"ERROR: {e}")
            raise

    out_path = os.path.join(args.output_dir, "template_context.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(ctx, f, indent=2)

    log(args.output_dir, f"output file written: {out_path}")
    print(f"Wrote {out_path}")
    print(json.dumps(ctx, indent=2))


if __name__ == "__main__":
    main()
